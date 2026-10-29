"""The "what to buy" scoreboard (pipeline/scoreboard.py), its /scoreboard text,
the brand warning after identification, and the /report sheet.

No eBay or Telegram calls; the DB is pointed at a temp file before import.
"""
import sys
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import db  # noqa: E402

db.DB_PATH = str(Path(tempfile.mkdtemp(prefix="scoreboard_")) / "test.db")
db.create_tables()

import telegram_bot as tb  # noqa: E402
from pipeline import scoreboard as sb  # noqa: E402

NOW = datetime(2026, 10, 29, tzinfo=timezone.utc)


def _iso(days_ago):
    return (NOW - timedelta(days=days_ago)).isoformat()


def _item(brand, item_type="tote bag", *, status="published", qty=1, cost=10.0,
          listed=90, sold_days=None, units=1, net=40.0, cat=None, shape="totals"):
    ebay = {"quantity": qty, "published_at": _iso(listed)}
    if cat:
        ebay["category_id"] = cat
    if sold_days is not None:
        if shape == "orders":
            ebay["orders"] = [{"units": 1, "net_proceeds": net / units,
                               "sold_at": _iso(sold_days - i)} for i in range(units)]
        else:
            ebay.update(sale_price=net + 10, net_proceeds=net, units_sold=units,
                        sold_at=_iso(sold_days))
    return {"item_id": f"{brand}-{id(ebay)}", "status": status, "created_at": _iso(listed),
            "identification": {"brand": brand, "item_type": item_type},
            "receipt": {"reduced_price": cost} if cost is not None else {},
            "listing": {"title": f"{brand} {item_type}"}, "ebay": ebay}


def _shop():
    items = []
    # Ed Hardy: 20 live, 4 sold of 24 → ~17%: pause
    items += [_item("Ed Hardy") for _ in range(20)]
    items += [_item("Ed Hardy", status="sold", qty=0, sold_days=60) for _ in range(4)]
    # Michael Kors aliases: 10 listings, 0 of 18 units, long exposure: stop
    for b in ("Michael Kors", "MK Access", "Michael Michael Kors", "Michael Kors Access"):
        items += [_item(b, "watch band", qty=2, listed=100) for _ in range(2)]
    items += [_item("Michael Kors", "watch band", qty=1, listed=100) for _ in range(2)]
    # Oakley: sells in a day, little exposure: still buy more
    items += [_item("Oakley", "sunglasses", status="sold", qty=0, listed=10 + i,
                    sold_days=9 + i) for i in range(5)]
    # Ralph Lauren aliases fold together
    items += [_item("Lauren Ralph Lauren", "duvet cover set", status="sold", qty=0,
                    sold_days=50), _item("Ralph Lauren", "duvet cover set")]
    # Two Bebe bags: too few to say anything
    items += [_item("Bebe") for _ in range(2)]
    # Never listed: ignored
    items += [{"item_id": "x", "status": "identified", "created_at": _iso(1),
               "identification": {"brand": "Ed Hardy"}}]
    return items


def _by(board, kind, name):
    return next(g for g in board[kind] if g["name"] == name)


def test_brand_aliases():
    assert sb.normalize_brand("Lauren Ralph Lauren") == "Ralph Lauren"
    assert sb.normalize_brand("Ralph Lauren Home") == "Ralph Lauren"
    assert sb.normalize_brand("MK Access") == "Michael Kors"
    assert sb.normalize_brand("Michael Michael Kors") == "Michael Kors"
    assert sb.normalize_brand("kate spade new york") == "Kate Spade"
    assert sb.normalize_brand(None) == "Unknown"
    assert sb.normalize_brand("  ") == "Unknown"


def test_verdicts():
    board = sb.build(_shop(), NOW)
    assert _by(board, "brands", "Ed Hardy")["verdict"] == "pause"
    mk = _by(board, "brands", "Michael Kors")
    assert mk["verdict"] == "stop" and mk["bought"] == 18 and mk["live"] == 10
    assert _by(board, "brands", "Oakley")["verdict"] == "buy more"
    assert _by(board, "categories", "Sunglasses")["verdict"] == "buy more"
    assert _by(board, "brands", "Bebe")["verdict"] == "ok"
    assert _by(board, "brands", "Ralph Lauren")["bought"] == 2
    assert board["total"]["listings"] == len(_shop()) - 1


def test_numbers():
    board = sb.build(_shop(), NOW)
    eh = _by(board, "brands", "Ed Hardy")
    assert (eh["live"], eh["bought"], eh["sold"]) == (20, 24, 4)
    assert abs(eh["sell_through"] - 4 / 24) < 1e-9
    assert eh["tied_up"] == 200.0                 # 20 live × $10
    assert eh["profit"] == 4 * (40.0 - 10.0)
    assert eh["median_days"] == 30.0              # listed 90d ago, sold 60d ago
    assert abs(eh["per100"] - 100 * 4 / (20 * 90 + 4 * 30)) < 1e-9


def test_per_order_list_and_missing_data():
    a = _item("Coach", "wallet", status="sold", qty=0, sold_days=20, units=3, net=90.0,
              shape="orders")
    f = sb.item_facts(a, NOW)
    assert f["sold"] == 3 and f["days_to_sell"] == 70.0 and f["profit"] == 90.0 - 30.0
    # ebay/orders.py's per-order shape: a "sales" list with units / net / sold_at
    b = {"status": "sold", "created_at": _iso(30), "receipt": {"reduced_price": 5},
         "identification": {"brand": "Coach"},
         "ebay": {"quantity": 0, "published_at": _iso(30), "units": 2, "net_proceeds": 70,
                  "sold_at": _iso(5), "first_sold_at": _iso(20),
                  "sales": [{"units": 1, "net": 30, "sold_at": _iso(20)},
                            {"units": 1, "net": 40, "sold_at": _iso(5)}]}}
    f = sb.item_facts(b, NOW)
    assert f["sold"] == 2 and f["profit"] == 60 and f["days_to_sell"] == 10.0
    del b["ebay"]["sales"]
    f = sb.item_facts(b, NOW)
    assert f["sold"] == 2 and f["profit"] == 60 and f["days_to_sell"] == 10.0
    bare = {"status": "sold", "created_at": _iso(5), "ebay": {}, "identification": None}
    f = sb.item_facts(bare, NOW)
    assert f["sold"] == 1 and f["profit"] is None and f["brand"] == "Unknown"
    nocost = _item("Coach", "wallet", cost=None, qty=2)
    f = sb.item_facts(nocost, NOW)
    assert f["tied_up"] == 0 and f["tied_up_unknown"] == 2
    assert sb.build([], NOW)["total"]["sell_through"] is None
    assert sb.format_scoreboard(sb.build([], NOW))


def test_category_from_id_type_and_title():
    assert sb.category_of({"ebay": {"category_id": "37644"}}) == "Duvets"
    assert sb.category_of({"identification": {"item_type": "Apple Watch Band"}}) == "Watch bands"
    assert sb.category_of({"identification": {"item_type": "Women's Wristwatch"}}) == "Watches"
    assert sb.category_of({"listing": {"title": "Kate Spade NY Band for Apple Watch"}}) == "Watch bands"
    adopted = {"status": "published", "identification": None, "created_at": _iso(5),
               "listing": {"title": "Kate Spade NY Black Band"}, "ebay": {"quantity": 1}}
    assert sb.brand_of(adopted, ["Kate Spade"]) == "Kate Spade"


def test_text_is_mobile_friendly():
    lines = sb.format_scoreboard(sb.build(_shop(), NOW))
    text = "\n".join(lines)
    assert "🟢 Buy more" in text and "🟡 Pause" in text and "🔴 Stop" in text
    assert "Michael Kors: 0/18 sold" in text
    assert all(len(l) <= 160 for l in lines)
    assert all(len(c) <= tb.TELEGRAM_MAX_CHARS for c in tb._chunk_lines(lines))


def test_saturation_warning():
    board = sb.build(_shop(), NOW)
    line = sb.saturation_warning({"identification": {"brand": "Ed Hardy"}}, board)
    assert line == "⚠️ Ed Hardy: 20 live, 17% sell-through — consider pausing."
    assert sb.saturation_warning({"identification": {"brand": "Oakley"}}, board) is None
    assert sb.saturation_warning({"identification": {"brand": "Nobody"}}, board) is None
    assert sb.saturation_warning({"identification": {}}, board) is None
    assert "Michael Kors" in sb.saturation_warning(
        {"identification": {"brand": "MK Access"}}, board)


def _fresh_db(monkeypatch):
    # Every test module points db.DB_PATH at its own file at import, so the
    # path in effect here may be another module's: use a clean one.
    monkeypatch.setattr(db, "DB_PATH", str(Path(tempfile.mkdtemp(prefix="scoreboard_")) / "t.db"))
    db.create_tables()


def test_identify_stage_adds_one_brand_line(monkeypatch):
    _fresh_db(monkeypatch)
    for it in _shop():
        if it["status"] in ("published", "sold") and it["identification"]["brand"] == "Ed Hardy":
            iid = db.create_item([])
            for field in ("identification", "receipt", "ebay", "listing"):
                db.update_field(iid, field, it[field])
            db.update_status(iid, it["status"])
    new = db.create_item([])
    db.update_field(new, "identification", {"brand": "ED HARDY", "item_type": "satchel"})
    lines = tb._brand_warnings(new)
    assert len(lines) == 1 and lines[0].startswith("⚠️ Ed Hardy: 20 live")
    other = db.create_item([])
    db.update_field(other, "identification", {"brand": "Oakley"})
    assert tb._brand_warnings(other) == []


def test_report_has_scoreboard_sheet(tmp_path, monkeypatch):
    _fresh_db(monkeypatch)
    import report
    import ebay.orders
    from openpyxl import load_workbook

    def no_ebay(_days):
        raise RuntimeError("no eBay in tests")
    monkeypatch.setattr(ebay.orders, "account_charges", no_ebay)
    out = report.build_report(str(tmp_path / "r.xlsx"))
    wb = load_workbook(out)
    assert "What to buy" in wb.sheetnames
    assert wb["What to buy"]["A1"].value.startswith("What to buy")
